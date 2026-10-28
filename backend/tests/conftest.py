"""Shared pytest fixtures.

Tests run against a **real PostgreSQL** database (never SQLite), using the
same tables and migrations as production. Every test is wrapped in an outer
transaction that is rolled back, so rows created by a test never survive it
and the dev database (including the seeded demo user) is left untouched.

Requires the database in ``DB_*`` / ``DATABASE_URL`` to be reachable; the
session fixture applies migrations first.
"""

import io
from collections.abc import AsyncIterator, Callable
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import (
    AsyncConnection,
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)
from sqlalchemy.pool import NullPool

from app.api.deps import get_db, get_storage
from app.config import settings
from app.main import app
from app.services.token_denylist import InMemoryTokenDenylist
from app.services.user_repository import PostgresUserRepository
from app.storage.local_storage import LocalStorage


@pytest.fixture(scope="session", autouse=True)
def migrated_database() -> None:
    """Apply migrations once so the ``users`` table exists."""
    command.upgrade(Config("alembic.ini"), "head")


@pytest.fixture(scope="session", autouse=True)
def postgres_user_store() -> None:
    """These tests exercise the real database, not the in-memory store."""
    settings.user_store = "postgres"


@pytest.fixture
async def db_connection() -> AsyncIterator[AsyncConnection]:
    """Open a connection whose transaction is rolled back after the test."""
    engine = create_async_engine(settings.database_dsn, poolclass=NullPool)
    async with engine.connect() as connection:
        transaction = await connection.begin()
        try:
            yield connection
        finally:
            await transaction.rollback()
    await engine.dispose()


@pytest.fixture
def session_factory(
    db_connection: AsyncConnection,
) -> async_sessionmaker[AsyncSession]:
    """Session factory bound to the test connection.

    ``create_savepoint`` turns the repository's ``commit()`` into a savepoint
    release, so service-level code works unchanged while the outer rollback
    still discards everything.
    """
    return async_sessionmaker(
        bind=db_connection,
        join_transaction_mode="create_savepoint",
        expire_on_commit=False,
    )


@pytest.fixture
async def db_session(
    session_factory: async_sessionmaker[AsyncSession],
) -> AsyncIterator[AsyncSession]:
    """Request-scoped session equivalent, isolated to the test."""
    async with session_factory() as session:
        yield session


@pytest.fixture
def user_repository(db_session: AsyncSession) -> PostgresUserRepository:
    """Repository wired to the test session."""
    return PostgresUserRepository(db_session)


@pytest.fixture
def denylist() -> InMemoryTokenDenylist:
    """Fresh revocation store per test."""
    return InMemoryTokenDenylist()


@pytest.fixture
def test_storage(tmp_path) -> LocalStorage:
    """Storage rooted in a temp directory so tests never touch uploads/."""
    return LocalStorage(tmp_path / "uploads")


@pytest.fixture
async def client(
    db_session: AsyncSession,
    test_storage: LocalStorage,
) -> AsyncIterator[AsyncClient]:
    """HTTP client for the app, sharing the test's transaction and storage."""

    async def override_get_db() -> AsyncIterator[AsyncSession]:
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_storage] = lambda: test_storage
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://testserver") as http:
        yield http
    app.dependency_overrides.clear()


@pytest.fixture
async def auth_headers(client: AsyncClient) -> dict[str, str]:
    """Register a fresh user and return a Bearer header for them."""
    response = await client.post(
        "/api/v1/auth/register",
        json={
            "name": "Resume Owner",
            "email": f"owner-{uuid4().hex[:12]}@example.com",
            "password": "Secret123",
        },
    )
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture
def email_factory() -> Callable[[str], str]:
    """Generate a unique email per call so tests never collide."""

    def make(prefix: str = "user") -> str:
        return f"{prefix}-{uuid4().hex[:12]}@example.com"

    return make


@pytest.fixture
def docx_bytes_factory() -> Callable[[str], bytes]:
    """Build a real DOCX (heading + paragraph + table) in memory."""

    def make(body: str) -> bytes:
        document = Document()
        document.add_heading("Experience", level=1)
        document.add_paragraph(body)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Python"
        table.rows[0].cells[1].text = "8 years"
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return make


@pytest.fixture
def pdf_bytes_factory() -> Callable[[str], bytes]:
    """Build a minimal, xref-correct single-page PDF with the given text."""

    def make(text: str) -> bytes:
        stream = f"BT /F1 10 Tf 50 750 Td 12 TL ({text}) Tj ET".encode("latin-1")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        out = bytearray(b"%PDF-1.4\n")
        offsets = []
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(out))
            out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
        xref_position = len(out)
        out += f"xref\n0 {len(objects) + 1}\n".encode()
        out += b"0000000000 65535 f \n"
        for offset in offsets:
            out += f"{offset:010d} 00000 n \n".encode()
        out += (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_position}\n%%EOF\n"
        ).encode()
        return bytes(out)

    return make
